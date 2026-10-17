"""Small fail-closed Excel formula evaluator; never executes Python or external links.

Supported: finite arithmetic, comparison, A1/range and sheet references,
SUM, SUMIF (aligned ranges), AVERAGE, MIN, MAX, COUNT, COUNTA, ABS, ROUND, IF.
Not an Excel replacement.
"""
from __future__ import annotations
import ast
from decimal import Decimal, ROUND_HALF_UP
import math
import operator
import re

class FormulaError(ValueError):
    pass

_REF = re.compile(r"(?<![A-Za-z0-9_\"'])((?:'[^']+'|[A-Za-z_][A-Za-z0-9_ ]*)!)?(\$?[A-Za-z]{1,3}\$?\d+)(?::(\$?[A-Za-z]{1,3}\$?\d+))?(?![A-Za-z0-9_])")


def evaluate_formula(formula: str, workbook=None, sheet: str | None = None, seen=None):
    if not isinstance(formula, str) or not formula.startswith("="):
        return formula
    if len(formula) > 8192 or any(token in formula for token in ["[", "]", "{", "}", "__", "http", "\\", ";"]):
        raise FormulaError("Unsupported or external formula")
    seen = set() if seen is None else seen
    refs = []
    def replace(match):
        sh = match.group(1)
        sh = sh[:-1].strip("'") if sh else sheet
        refs.append((sh, match.group(2).replace("$", ""), (match.group(3) or "").replace("$", "")))
        return f"REF({len(refs)-1})"
    # Excel strings (including criteria like ">=5" and labels like "A1")
    # must not be rewritten as operators or cell references.
    parts = re.split(r'("(?:[^"]|"")*")', formula[1:])
    for i, part in enumerate(parts):
        if i % 2:
            parts[i] = repr(part[1:-1].replace('""', '"'))
        else:
            part = part.replace("^", "**").replace("<>", "!=")
            part = re.sub(r"(?<![<>=!])=(?!=)", "==", part)
            part = re.sub(r"(\d+(?:\.\d+)?)%", r"(\1/100)", part)
            parts[i] = _REF.sub(replace, part)
    expr = "".join(parts)
    try:
        tree = ast.parse(expr, mode="eval")
    except (SyntaxError, RecursionError) as exc:
        raise FormulaError("Invalid formula") from exc
    if len(list(ast.walk(tree))) > 4096:
        raise FormulaError("Formula complexity limit")
    def finite(value):
        if isinstance(value, (int, float)) and (not math.isfinite(value) or abs(value) > 1e15):
            raise FormulaError("Nonfinite or out-of-bounds result")
        return value
    def flatten(items):
        result = []
        for item in items:
            result.extend(flatten(item) if isinstance(item, list) else [item])
        return result
    def numeric(args):
        return [x for x in flatten(args) if isinstance(x, (int, float)) and not isinstance(x, bool)]
    def ref(index, preserve_blanks=False):
        if workbook is None:
            raise FormulaError("Cell reference requires workbook")
        sh, start, end = refs[index]
        if sh not in workbook.sheetnames:
            raise FormulaError("Unknown sheet")
        ws = workbook[sh]
        def cell_value(address):
            key = (sh, address)
            if key in seen:
                raise FormulaError("Cyclic formula")
            value = ws[address].value
            if isinstance(value, str) and value.startswith("="):
                return evaluate_formula(value, workbook, sh, seen | {key})
            return 0 if value is None and not preserve_blanks else value
        if end:
            from openpyxl.utils.cell import range_boundaries
            bounds = range_boundaries(f"{start}:{end}")
            if (bounds[2]-bounds[0]+1)*(bounds[3]-bounds[1]+1) > 50000:
                raise FormulaError("Range too large")
            return [cell_value(c.coordinate) for row in ws[f"{start}:{end}"] for c in row]
        return cell_value(start)
    def raw_argument(node):
        if (isinstance(node, ast.Call) and isinstance(node.func, ast.Name)
                and node.func.id == "REF" and len(node.args) == 1
                and isinstance(node.args[0], ast.Constant)):
            return ref(node.args[0].value, preserve_blanks=True)
        return walk(node)
    def matches(value, criterion):
        if criterion is None:
            criterion = 0
        op, target = "=", criterion
        if isinstance(criterion, str):
            match = re.match(r"^(<=|>=|<>|=|<|>)(.*)$", criterion, re.S)
            if match:
                op, target = match.groups()
            try:
                target = float(target)
            except ValueError:
                pass
        if isinstance(target, (int, float)) and not isinstance(target, bool):
            if value is None:
                value = 0
            if isinstance(value, bool) or not isinstance(value, (int, float)):
                return op == "<>"
        elif isinstance(target, str):
            if value is None:
                value = ""
            if not isinstance(value, str):
                return op == "<>"
            value, target = value.casefold(), target.casefold()
            if op in {"=", "<>"}:
                pattern, i = [], 0
                while i < len(target):
                    char = target[i]
                    if char == "~" and i + 1 < len(target):
                        i += 1
                        pattern.append(re.escape(target[i]))
                    else:
                        pattern.append(".*" if char == "*" else "." if char == "?" else re.escape(char))
                    i += 1
                equal = re.fullmatch("".join(pattern), value, re.S) is not None
                return equal if op == "=" else not equal
        return {"=": operator.eq, "<>": operator.ne, "<": operator.lt,
                "<=": operator.le, ">": operator.gt, ">=": operator.ge}[op](value, target)
    def walk(node):
        if isinstance(node, ast.Expression): return walk(node.body)
        if isinstance(node, ast.Constant) and isinstance(node.value, (int, float, bool, str)):
            return finite(node.value)
        if isinstance(node, ast.Name) and node.id in {"TRUE", "FALSE"}: return node.id == "TRUE"
        if isinstance(node, ast.UnaryOp) and isinstance(node.op, (ast.UAdd, ast.USub)):
            value = walk(node.operand)
            return finite(value if isinstance(node.op, ast.UAdd) else -value)
        if isinstance(node, ast.BinOp):
            ops = {ast.Add: operator.add, ast.Sub: operator.sub, ast.Mult: operator.mul, ast.Div: operator.truediv, ast.Pow: operator.pow}
            if type(node.op) not in ops: raise FormulaError("Unsupported operator")
            left, right = walk(node.left), walk(node.right)
            if not isinstance(left, (int, float)) or not isinstance(right, (int, float)): raise FormulaError("Nonnumeric arithmetic")
            if isinstance(node.op, ast.Pow) and abs(right) > 20: raise FormulaError("Exponent bound")
            try: return finite(ops[type(node.op)](left, right))
            except (ArithmeticError, ValueError, TypeError) as exc: raise FormulaError(str(exc)) from exc
        if isinstance(node, ast.Compare) and len(node.ops) == 1:
            op = {ast.Eq: operator.eq, ast.NotEq: operator.ne, ast.Lt: operator.lt, ast.LtE: operator.le, ast.Gt: operator.gt, ast.GtE: operator.ge}.get(type(node.ops[0]))
            if not op: raise FormulaError("Unsupported comparison")
            return op(walk(node.left), walk(node.comparators[0]))
        if isinstance(node, ast.Call) and isinstance(node.func, ast.Name) and not node.keywords:
            name = node.func.id.upper()
            if name == "IF" and len(node.args) == 3: return walk(node.args[1] if walk(node.args[0]) else node.args[2])
            if name in {"COUNT", "COUNTA"}:
                values = flatten([raw_argument(x) for x in node.args])
                return len(numeric(values)) if name == "COUNT" else sum(x is not None for x in values)
            if name == "SUMIF":
                if len(node.args) not in {2, 3}:
                    raise FormulaError("SUMIF requires two or three arguments")
                selected = flatten([raw_argument(node.args[0])])
                criterion = walk(node.args[1])
                amounts = flatten([raw_argument(node.args[2])]) if len(node.args) == 3 else selected
                if len(selected) != len(amounts):
                    raise FormulaError("SUMIF requires aligned ranges")
                return finite(sum(amount for value, amount in zip(selected, amounts)
                                  if isinstance(amount, (int, float)) and not isinstance(amount, bool)
                                  and matches(value, criterion)))
            args = [walk(x) for x in node.args]
            if name == "REF" and len(args) == 1 and isinstance(args[0], int): return ref(args[0])
            values = numeric(args)
            funcs = {"SUM": lambda: sum(values), "AVERAGE": lambda: sum(values)/len(values), "MIN": lambda: min(values), "MAX": lambda: max(values), "COUNT": lambda: len(values), "ABS": lambda: abs(args[0]), "ROUND": lambda: float(Decimal(str(args[0])).quantize(Decimal(1).scaleb(-(int(args[1]) if len(args)>1 else 0)), rounding=ROUND_HALF_UP))}
            if name not in funcs: raise FormulaError(f"Unsupported function {name}")
            try: return finite(funcs[name]())
            except (ArithmeticError, TypeError, ValueError, IndexError) as exc: raise FormulaError(str(exc)) from exc
        raise FormulaError("Unsupported formula syntax")
    return walk(tree)
