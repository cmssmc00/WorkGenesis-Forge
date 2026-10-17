"""Real office/PDF artifact IO with deterministic packages and bounded extraction."""
from __future__ import annotations
import csv
import io
import json
import re
import zipfile
from datetime import datetime
from pathlib import Path
from xml.sax.saxutils import escape
import textwrap
from .safeformula import evaluate_formula

MAX_BYTES = 20_000_000
MAX_CELLS = 100_000
FIXED_TIME = datetime(2025, 1, 1)


def normalize_zip(path: Path):
    with zipfile.ZipFile(path) as src:
        contents = {n: src.read(n) for n in src.namelist()}
    if "docProps/core.xml" in contents:
        xml=contents["docProps/core.xml"].decode("utf-8")
        xml=re.sub(r"(<dcterms:(?:created|modified)[^>]*>)[^<]*(</dcterms:(?:created|modified)>)",r"\g<1>2025-01-01T00:00:00Z\g<2>",xml)
        contents["docProps/core.xml"]=xml.encode("utf-8")
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=6) as dest:
        for name in sorted(contents):
            info = zipfile.ZipInfo(name, (2025, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            dest.writestr(info, contents[name])
    path.write_bytes(out.getvalue())


def _tables(spec):
    tables = spec.get("tables", [])
    if not isinstance(tables, list): raise ValueError("tables must be a list")
    count = 0
    for table in tables:
        headers, rows = table.get("headers", []), table.get("rows", [])
        if not headers or len(headers) > 100: raise ValueError("Invalid table headers")
        count += sum(len(row) for row in rows)
        if any(len(row) != len(headers) for row in rows): raise ValueError("Table width mismatch")
    if count > MAX_CELLS: raise ValueError("Too many cells")
    return tables


def write_artifact(path: Path, spec: dict) -> None:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    if len(json.dumps(spec)) > MAX_BYTES: raise ValueError("Artifact input too large")
    title = str(spec.get("title", path.stem))
    text = str(spec.get("text", ""))
    tables = _tables(spec)
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md", ".eml", ".html", ".ics"}:
        parts = [title, text]
        for tab in tables:
            parts.extend([tab.get("name", "Table"), "\t".join(map(str, tab["headers"]))])
            parts.extend("\t".join(map(str, row)) for row in tab["rows"])
        path.write_text("\n\n".join(parts)+"\n", encoding="utf-8")
    elif suffix == ".json":
        path.write_text(json.dumps(spec, indent=2, ensure_ascii=False)+"\n", encoding="utf-8")
    elif suffix == ".csv":
        if len(tables) != 1: raise ValueError("CSV requires exactly one table")
        for row in [tables[0]["headers"], *tables[0]["rows"]]:
            for cell in row:
                if isinstance(cell, str) and cell.lstrip().startswith(("=", "+", "@", "-")):
                    try: float(cell)
                    except ValueError: raise ValueError("CSV formula-like strings are forbidden; use literal numeric values")
        with path.open("w", newline="", encoding="utf-8") as f:
            w = csv.writer(f); w.writerow(tables[0]["headers"]); w.writerows(tables[0]["rows"])
    elif suffix == ".xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
        from openpyxl.workbook.properties import CalcProperties
        wb = Workbook(); wb.remove(wb.active)
        wb.properties.created = FIXED_TIME; wb.properties.modified = FIXED_TIME
        wb.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True)
        xlsx_tables=list(tables)
        if text: xlsx_tables.append({"name":"Source Notes","headers":["context"],"rows":[[title],*[ [line] for line in text.splitlines() if line.strip() ]]})
        for tab in xlsx_tables or [{"name":"Notes","headers":["text"],"rows":[[text]]}]:
            name = re.sub(r"[\[\]:*?/\\]", "_", str(tab.get("name", "Data")))[:31]
            ws = wb.create_sheet(name); ws.append(tab["headers"])
            for row in tab["rows"]: ws.append(row)
            ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
            ws.sheet_properties.pageSetUpPr.fitToPage=True
            ws.page_setup.orientation="landscape"; ws.page_setup.paperSize=ws.PAPERSIZE_A3
            ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
            ws.print_title_rows="1:1"
            for c in ws[1]: c.fill = PatternFill("solid", fgColor="183B56"); c.font = Font(color="FFFFFF", bold=True)
            for col in range(1, len(tab["headers"])+1):
                samples = [str(tab["headers"][col-1])] + [str(row[col-1] if row[col-1] is not None else "") for row in tab["rows"][:200]]
                ws.column_dimensions[get_column_letter(col)].width = min(70, max(16, max(len(value) for value in samples) + 3))
            for row in ws.iter_rows(min_row=2):
                for c in row:
                    c.alignment = Alignment(vertical="top", wrap_text=True)
                    if c.column == 3: c.number_format = "#,##0.00"
            if len(tab["headers"]) <= 6:
                ws.page_setup.paperSize = ws.PAPERSIZE_A4
            if name == "Summary" and len(tab["headers"]) == 2:
                ws.column_dimensions["A"].width = 31
                ws.column_dimensions["B"].width = 100
                for r in range(2, ws.max_row + 1):
                    value_length = len(str(ws.cell(r, 2).value or ""))
                    ws.row_dimensions[r].height = max(24, 16 * ((value_length // 95) + 1))
            if name=="Metrics":
                for letter,width in {"A":28,"B":36,"C":26,"D":18,"E":28,"F":80}.items():ws.column_dimensions[letter].width=width
                for r in range(2,ws.max_row+1):ws.row_dimensions[r].height=64
        wb.save(path); normalize_zip(path)
    elif suffix == ".docx":
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor
        doc = Document(); props = doc.core_properties
        props.created = FIXED_TIME; props.modified = FIXED_TIME; props.author = "" if spec.get("business_document") else "WorkGenesis Forge"
        doc.styles["Normal"].font.name="Calibri"; doc.styles["Normal"].font.size=Pt(10)
        sec=doc.sections[0]; sec.header.paragraphs[0].text="" if spec.get("business_document") else "WORKGENESIS FORGE • SCENARIO WORKING DOCUMENT"
        doc.add_heading(title,0)
        for para in text.split("\n"):
            if para.strip():
                if para.strip() in {"Decision","Evidence","Risks","Next steps"}:doc.add_heading(para.strip(),1)
                else:doc.add_paragraph(para.strip())
        for tab in tables:
            doc.add_heading(tab.get("name","Evidence"),2)
            tbl=doc.add_table(rows=1,cols=len(tab["headers"])); tbl.style="Light Shading Accent 1"
            for i,v in enumerate(tab["headers"]): tbl.rows[0].cells[i].text=str(v)
            for row in tab["rows"]:
                cells=tbl.add_row().cells
                for i,v in enumerate(row): cells[i].text=str(v if v is not None else "")
        doc.save(path); normalize_zip(path)
    elif suffix == ".pptx":
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor
        prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
        prs.core_properties.created=FIXED_TIME; prs.core_properties.modified=FIXED_TIME
        slides=[dict(item) for item in spec.get("slides", [])] or [{"title":title,"text":text}]
        # Tables are paginated into body text to keep long evidence readable.
        for tab in tables:
            for start in range(0,len(tab["rows"]),5):
                lines=[" | ".join(str(x) for x in row) for row in tab["rows"][start:start+5]]
                slides.append({"title":tab.get("name","Evidence"),"text":" | ".join(map(str,tab["headers"]))+"\n"+"\n\n".join(lines)})
        for data in slides:
            # Keep slides bounded: long paragraphs get continuation slides.
            content=str(data.get("text", ""));chunks=[];current=[];size=0
            for line in content.splitlines():
                # Preserve complete claims and paragraphs; unusually long lines wrap at words.
                pieces=textwrap.wrap(line,width=1150,break_long_words=False,break_on_hyphens=False) or [""]
                for piece in pieces:
                    if current and size+len(piece)>1300:chunks.append("\n".join(current));current=[];size=0
                    current.append(piece);size+=len(piece)+1
            if current:chunks.append("\n".join(current))
            chunks=chunks or [""]
            for part,chunk in enumerate(chunks):
                slide=prs.slides.add_slide(prs.slide_layouts[6])
                bg=slide.background.fill; bg.solid(); bg.fore_color.rgb=RGBColor.from_string("F7F9FC")
                box=slide.shapes.add_textbox(Inches(.6),Inches(.35),Inches(12.1),Inches(.8))
                p=box.text_frame.paragraphs[0]; p.text=str(data.get("title",title))+(" (continued)" if part else ""); p.font.size=Pt(28); p.font.bold=True; p.font.color.rgb=RGBColor.from_string("183B56")
                body=slide.shapes.add_textbox(Inches(.65),Inches(1.4),Inches(12),Inches(5.6))
                body.text_frame.word_wrap=True
                for idx,line in enumerate(chunk.splitlines()):
                    p=body.text_frame.paragraphs[0] if idx==0 else body.text_frame.add_paragraph(); p.text=line; p.font.size=Pt(17)
        prs.save(path); normalize_zip(path)
    elif suffix == ".pdf":
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, LongTable, TableStyle
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import landscape,A4
        from reportlab.pdfgen import canvas
        class DeterministicCanvas(canvas.Canvas):
            def __init__(self,*args,**kwargs): kwargs["invariant"]=1; super().__init__(*args,**kwargs)
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        import reportlab
        font_root=Path(reportlab.__file__).parent/"fonts"
        for font_name,file_name in [("ForgeVera","Vera.ttf"),("ForgeVeraBold","VeraBd.ttf"),("ForgeVeraItalic","VeraIt.ttf"),("ForgeVeraBoldItalic","VeraBI.ttf")]:
            if font_name not in pdfmetrics.getRegisteredFontNames():pdfmetrics.registerFont(TTFont(font_name,str(font_root/file_name)))
        pdfmetrics.registerFontFamily("ForgeVera",normal="ForgeVera",bold="ForgeVeraBold",italic="ForgeVeraItalic",boldItalic="ForgeVeraBoldItalic")
        styles=getSampleStyleSheet()
        for style in styles.byName.values():style.fontName="ForgeVeraBold" if style.name in {"Title","Heading1","Heading2","Heading3"} else "ForgeVera"
        styles["BodyText"].fontSize=9;styles["BodyText"].leading=12
        story=[Paragraph(escape(title),styles["Title"]),Spacer(1,12)]
        for para in text.splitlines():
            if para.strip(): story += [Paragraph(escape(para),styles["BodyText"]),Spacer(1,5)]
        for tab in tables:
            story += [Spacer(1,10),Paragraph(escape(tab.get("name","Evidence")),styles["Heading2"])]
            rows=[[Paragraph(escape(str(x if x is not None else "")),styles["BodyText"]) for x in row] for row in [tab["headers"],*tab["rows"]]]
            t=LongTable(rows,repeatRows=1,colWidths=[740/len(tab["headers"])]*len(tab["headers"]))
            t.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#DBE7F3")),("VALIGN",(0,0),(-1,-1),"TOP"),("GRID",(0,0),(-1,-1),.3,colors.lightgrey),("FONTSIZE",(0,0),(-1,-1),8)]))
            story.append(t)
        doc=SimpleDocTemplate(str(path),pagesize=landscape(A4),leftMargin=45,rightMargin=45,topMargin=36,bottomMargin=36)
        doc.build(story,canvasmaker=DeterministicCanvas)
    else:
        raise ValueError(f"Unsupported artifact format: {suffix}")


def _check_file(path):
    if path.is_symlink() or not path.is_file(): raise ValueError("Missing file or symlink")
    if path.stat().st_size > MAX_BYTES: raise ValueError("Artifact too large")
    if path.suffix.lower() in {".xlsx", ".docx", ".pptx"}:
        with zipfile.ZipFile(path) as z:
            if sum(i.file_size for i in z.infolist()) > 100_000_000: raise ValueError("Expanded archive too large")
            if any("vbaproject" in i.filename.lower() or "externallink" in i.filename.lower() for i in z.infolist()): raise ValueError("Macros/external links forbidden")


def extract_artifact(path: Path) -> dict:
    path=Path(path); _check_file(path); suffix=path.suffix.lower(); result={"text":"","tables":[],"formulas":[],"format":suffix[1:],"parser_notes":[],"unverified_visual_content":False}
    if suffix in {".png", ".jpg", ".jpeg", ".webp"}:
        from PIL import Image
        with Image.open(path) as picture:
            picture.verify()
        with Image.open(path) as picture:
            result["image_size"] = list(picture.size)
            result["text"] = f"Image {path.name}: {picture.width} x {picture.height} pixels. Pixel content requires view_image."
        result["unverified_visual_content"] = True
        result["parser_notes"].append("Image decoding verified; no visual content inferred from metadata.")
    elif suffix in {".wav", ".mp3", ".mp4", ".mov"}:
        result["text"] = f"Media attachment {path.name}, {path.stat().st_size} bytes. Decode and inspect with code_exec; metadata is not content evidence."
        result["unverified_visual_content"] = True
        result["unverified_media_content"] = True
        result["parser_notes"].append("No transcript or frame evidence supplied by this host parser; tool-based inspection is required.")
        if suffix == ".wav":
            import wave
            with wave.open(str(path), "rb") as audio:
                result["media_metadata"] = {"channels": audio.getnchannels(), "sample_rate": audio.getframerate(),
                                            "frames": audio.getnframes()}
    elif suffix in {".txt",".md",".eml",".html",".ics"}: result["text"]=path.read_text(encoding="utf-8")
    elif suffix==".json": result["text"]=path.read_text(encoding="utf-8"); json.loads(result["text"])
    elif suffix==".csv":
        with path.open(newline="",encoding="utf-8-sig") as f: rows=list(csv.reader(f))
        result["tables"]=[{"name":path.stem,"headers":rows[0] if rows else [],"rows":rows[1:]}]
    elif suffix==".xlsx":
        from openpyxl import load_workbook
        with zipfile.ZipFile(path) as z:
            result["unverified_visual_content"] = any(n.startswith(("xl/media/", "xl/charts/", "xl/drawings/", "xl/embeddings/")) for n in z.namelist())
        result["parser_notes"].append("Workbook cells and bounded formulas extracted; chart/image/layout content is not visually verified.")
        wb=load_workbook(path,data_only=False,read_only=False,keep_links=False)
        if sum(ws.max_row*ws.max_column for ws in wb)>MAX_CELLS: raise ValueError("Too many workbook cells")
        for ws in wb:
            rows=[]
            for row in ws.iter_rows():
                values=[]
                for cell in row:
                    value=cell.value
                    if isinstance(value,str) and value.startswith("="):
                        value=evaluate_formula(value,wb,ws.title,{(ws.title,cell.coordinate)})
                        result["formulas"].append({"sheet":ws.title,"cell":cell.coordinate,"formula":cell.value,"value":value})
                    if isinstance(value,datetime): value=value.isoformat()
                    values.append(value)
                rows.append(values)
            result["tables"].append({"name":ws.title,"headers":rows[0] if rows else [],"rows":rows[1:]})
        wb.close()
    elif suffix==".docx":
        from docx import Document
        doc=Document(path)
        from xml.etree import ElementTree as ET
        with zipfile.ZipFile(path) as z:
            parts=[n for n in z.namelist() if n.startswith("word/") and n.endswith(".xml") and any(k in n for k in ["document", "header", "footer", "footnotes", "endnotes", "comments"])]
            texts=[]
            for part in parts:
                xml=ET.fromstring(z.read(part))
                for para in xml.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p"):
                    texts.append("".join(t.text or "" for t in para.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t")))
            result["text"]="\n".join(texts)
            result["unverified_visual_content"]=any(n.startswith("word/media/") or n.startswith("word/embeddings/") for n in z.namelist())
            result["parser_notes"].append("All text paragraphs in body, tables, headers, footers, footnotes, endnotes and comments extracted; no raster/OLE interpretation.")
        for idx,tab in enumerate(doc.tables):
            rows=[[c.text for c in row.cells] for row in tab.rows]
            result["tables"].append({"name":f"Table{idx+1}","headers":rows[0] if rows else [],"rows":rows[1:]})
    elif suffix==".pptx":
        from pptx import Presentation
        prs=Presentation(path)
        from xml.etree import ElementTree as ET
        with zipfile.ZipFile(path) as z:
            parts=[n for n in z.namelist() if n.endswith(".xml") and (re.fullmatch(r"ppt/slides/slide\d+\.xml",n) or re.fullmatch(r"ppt/notesSlides/notesSlide\d+\.xml",n))]
            texts=[]
            for part in sorted(parts):
                xml=ET.fromstring(z.read(part))
                for paragraph in xml.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}p"):
                    texts.append("".join(t.text or "" for t in paragraph.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}t")))
            result["text"]="\n".join(texts); result["slides"]=len(prs.slides)
            result["unverified_visual_content"]=any(n.startswith("ppt/media/") or n.startswith("ppt/embeddings/") for n in z.namelist())
            result["parser_notes"].append("Grouped text, slide notes and table text extracted; image-only/OLE content requires separate review.")
    elif suffix==".pdf":
        from pypdf import PdfReader
        reader=PdfReader(path)
        if reader.is_encrypted or len(reader.pages)>500: raise ValueError("Encrypted or oversized PDF")
        result["text"]="\n".join(p.extract_text() or "" for p in reader.pages); result["pages"]=len(reader.pages)
        result["unverified_visual_content"]=any(len(p.images)>0 for p in reader.pages)
        result["parser_notes"].append("PDF embedded text extracted; scanned images/OCR and geometric layout grading are unsupported.")
    else: raise ValueError(f"Unsupported format: {suffix}")
    table_text="\n".join("\n".join(" | ".join(str(x if x is not None else "") for x in row) for row in [tab["headers"],*tab["rows"]]) for tab in result["tables"])
    if table_text: result["text"]+="\n"+table_text
    return result
