"""Regression cases for formulas found in real interrupted companion files."""
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from wgforge.artifacts import extract_artifact
from wgforge.safeformula import FormulaError, evaluate_formula


class SafeFormulaTests(unittest.TestCase):
    def setUp(self):
        self.wb = Workbook()
        self.ws = self.wb.active
        self.ws.title = "Costs"
        for row in [("Type", "Amount"), ("One-time", 10), ("Annual", 20),
                    ("one-TIME", 30), (None, None), ("A1", 5), ("A*", 7)]:
            self.ws.append(row)

    def evaluate(self, formula):
        return evaluate_formula(formula, self.wb, "Costs")

    def test_sumif_text_case_quotes_and_cross_sheet(self):
        cases = {
            '=SUMIF(A2:A7,"One-time",B2:B7)': 40,
            '=SUMIF(A2:A7,"Annual",B2:B7)': 20,
            '=SUMIF(A2:A7,"A1",B2:B7)': 5,
            '=SUMIF(A2:A7,"A*",B2:B7)': 32,
            '=SUMIF(A2:A7,"A~*",B2:B7)': 7,
            '=SUMIF(A2:A7,"<>Annual",B2:B7)': 52,
            '=SUMIF(\'Costs\'!$A$2:$A$7,"One-time",\'Costs\'!$B$2:$B$7)': 40,
        }
        for formula, expected in cases.items():
            with self.subTest(formula=formula):
                self.assertEqual(self.evaluate(formula), expected)
        self.ws["D1"] = 'He said "yes"'
        self.assertEqual(self.evaluate('=SUMIF(D1,"He said ""yes""",B2)'), 10)

    def test_sumif_numeric_criteria_and_omitted_sum_range(self):
        for criterion, expected in [('>=20', 50), ('<10', 12), ('<>20', 52), ('=10', 10)]:
            with self.subTest(criterion=criterion):
                self.assertEqual(self.evaluate(f'=SUMIF(B2:B7,"{criterion}")'), expected)
        self.ws["D1"] = ">=20"
        self.assertEqual(self.evaluate('=SUMIF(B2:B7,D1)'), 50)
        self.assertEqual(self.evaluate('=SUMIF(B2:B7,10)'), 10)

    def test_counta_distinguishes_blank_zero_boolean_and_formula_empty_string(self):
        for i, value in enumerate([None, 0, False, "text", '=""', '=1+2'], 1):
            self.ws.cell(i, 5, value)
        self.assertEqual(self.evaluate('=COUNTA(E1:E6)'), 5)
        self.assertEqual(self.evaluate('=COUNT(E1:E6)'), 2)
        self.assertEqual(self.evaluate('=COUNTA(E1,E2,"",FALSE)'), 3)
        self.assertEqual(self.evaluate('=E1+2'), 2)

    def test_native_workbook_extraction(self):
        self.ws["D2"] = '=SUMIF(A2:A7,"One-time",B2:B7)'
        self.ws["D3"] = '=COUNTA(A2:A7)'
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "intake.xlsx"
            self.wb.save(path)
            parsed = extract_artifact(path)
        self.assertEqual([item["value"] for item in parsed["formulas"]], [40, 5])

    def test_invalid_ranges_cycles_and_external_formulas_still_fail(self):
        for formula in ['=SUMIF(A2:A7,"Annual",B2:B3)', '=SUMIF(A2:A7)',
                        '=WEBSERVICE("http://example.com")', '=SUM([external.xlsx]Sheet1!A1)',
                        '=UNKNOWN(1)', '=__import__("os")']:
            with self.subTest(formula=formula), self.assertRaises(FormulaError):
                self.evaluate(formula)
        self.ws["D1"] = '=SUMIF(A2,"One-time",D1)'
        with self.assertRaisesRegex(FormulaError, "Cyclic"):
            self.evaluate('=D1')


if __name__ == "__main__":
    unittest.main()
