"""Full author/execute/audit/attribute/repair/reexecute HTTP protocol with fake transport.

Uses genuine saved USGS anchor bytes, real native writers and actual orchestration.
All provider messages are deterministic QA fixtures, never genuine model evidence.
"""
import sys,json,os,tempfile,copy,csv,shutil
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch
import unittest
from _qa_paths import REPO, QA_OUTPUT, TOKENIZER
sys.path.insert(0,str(REPO))
from wgforge import workflow as wf,runtime as rt
from wgforge.common import read_json,write_json,ForgeError
from wgforge.materialize import verify_bundle,inspect_outputs
from wgforge.schema import VERSION
from wgforge.dataset import export_sft,qualify,sign_review
ROOT=REPO
QA=QA_OUTPUT
PRIVATE='PRIVATE_RUBRIC_CANARY_a630f2'
REPAIR='PRIVATE_REPAIR_NOTES_CANARY_19bdf2'
class Response:
 def __init__(self,obj):self.raw=json.dumps(obj).encode()
 def __enter__(self):return self
 def __exit__(self,*a):pass
 def read(self,n):return self.raw[:n]
def response(message):return {'id':'QA_MOCK_RESPONSE','choices':[{'message':message,'finish_reason':'tool_calls' if message.get('tool_calls') else 'stop'}],'usage':{'prompt_tokens':100,'completion_tokens':100}}
def json_reply(result):return response({'role':'assistant','content':json.dumps(result)})
def tool(name,args,ident):return response({'role':'assistant','content':None,'tool_calls':[{'id':ident,'type':'function','function':{'name':name,'arguments':json.dumps(args)}}]})
class Transport:
 def __init__(self,queues):self.queues=queues;self.requests=[]
 def open(self,request,timeout=None):
  payload=json.loads(request.data);self.requests.append(payload);role=payload['model'].replace('QA_MOCK_ONLY_','')
  if not self.queues[role]:raise AssertionError('Unexpected mocked provider request: '+role)
  result=self.queues[role].pop(0)
  if callable(result):result=result(payload)
  return Response(result)
def cfg():
 c={'approve_live':True,'endpoint':'https://qa.example.org/v1/chat/completions','allowed_hosts':['qa.example.org'],'api_key_env':'WG_QA_ONLY_KEY','privacy':{'public_upload_approved':True,'private_upload_approved':True},'timeout_seconds':2,'max_retries':0,'budget':{'max_cost_usd':5,'input_usd_per_million':1,'output_usd_per_million':1,'max_requests':20,'max_steps':20,'max_input_tokens':1000000,'max_output_tokens':10000,'max_total_tokens':2000000}}
 return {'approve_live':True,'campaign_budget':{'max_cost_usd':50,'max_requests':100,'max_tokens':10000000},'roles':{r:{**copy.deepcopy(c),'model':'QA_MOCK_ONLY_'+r} for r in ['author','solver','judge','attributor','repairer','repair_reviewer','teacher']},'repair':{'max_rounds':6,'repeated_defect_limit':3}}
def setup_protocol():
 manifest_path=ROOT/'examples/anchors/usgs_manifest.json';manifest=read_json(manifest_path);source=manifest['sources'][0]
 with (manifest_path.parent/source['path']).open(newline='') as f:rows=list(csv.DictReader(f))
 expected={'count':len(rows),'threshold':sum(Decimal(r['mag'])>=Decimal('4.5') for r in rows if r['mag'])}
 doc={'schema_version':VERSION,'scenario_id':'overwritten','lineage_id':'overwritten','revision':0,'anchor_sources':manifest['sources'],
 'request':{'title':'Catalog monitoring QA task','role':'Geoscientist','context':'Retrospective bounded USGS catalog pilot for a fictional team.','instructions':'Count all rows of the bounded USGS snapshot and state its source coverage. Produce the requested workbook and memo.','deliverables':[{'path':'analysis.xlsx','format':'xlsx','description':'Summary metrics from USGS rows'},{'path':'memo.docx','format':'docx','description':'Explain counts and bounded source scope'}]},
 'companions':[{'id':'scope','path':'scope.md','format':'md','fictitious':True,'disclosure':'Fictional internal planning memo for this exercise.','spec':{'title':'Internal scope','text':'Use this bounded catalog as a retrospective analytical sample; do not make safety predictions.'}}],
 'rubric':{'items':[{'id':'count','weight':5,'criterion':'Workbook count equals source row count '+PRIVATE,'evidence':['USGS source CSV'],'verification':'structured'},{'id':'threshold','weight':5,'criterion':'Workbook counts events with magnitude at least 4.5','evidence':['USGS mag column'],'verification':'structured'},{'id':'scope','weight':3,'criterion':'Memo identifies bounded 200-event sample','evidence':['request scope'],'verification':'text'}]},'private_reference':{'expected':expected,'private_note':PRIVATE}}
 def author(p):
  body=json.loads(p['messages'][1]['content']);source=body['sources'][0]
  return json_reply({'source_assessment':[{'source_id':source['record']['id'],'usable':True,'occupation_fit':'A geoscientist can use this catalog to count and compare historical event magnitudes.','professional_operations':['Aggregate source rows and compare magnitude thresholds'], 'evidence_excerpt':source['extracted_text'][:80], 'limitations':'Bounded sample only; not a full earthquake census.'}], 'scenario':doc})
 def repair(p):
  body=json.loads(p['messages'][1]['content']);updated=copy.deepcopy(body['scenario']);updated['revision']+=1;updated['request']['instructions']+=' Also count events with magnitude at least 4.5.';updated['private_reference']['repair_note']=REPAIR
  return json_reply({'scenario':updated,'changes':[{'criterion_id':'threshold','reason':'Initial request omitted the required threshold','changed_paths':['request.instructions']}]})
 def solve(seq,threshold):
  metrics=[['event_count',expected['count']]]+([['threshold_count',expected['threshold']]] if threshold else [])
  return [tool('read_file',{'path':'input/anchors/'+source['id']+'.csv'},seq+'read'),tool('write_workbook',{'path':'output/analysis.xlsx','spec':{'title':'Catalog metrics','tables':[{'name':'Metrics','headers':['metric','value'],'rows':metrics}]}},seq+'xlsx'),tool('write_document',{'path':'output/memo.docx','spec':{'title':'Retrospective catalog memo','text':f'The bounded USGS source contains {expected["count"]} events and is not a complete census.'+(f' There are {expected["threshold"]} events with magnitude at least 4.5.' if threshold else '')}},seq+'docx'),tool('finish',{'summary':'QA fixture files written'},seq+'finish')]
 def grade(full):return json_reply({'scores':[{'criterion_id':'count','score':5,'rationale':'Actual workbook count matches source','evidence':['analysis.xlsx Metrics event_count']},{'criterion_id':'threshold','score':5 if full else 0,'rationale':'Threshold count matches source' if full else 'Threshold was not specified in request and not present','evidence':['analysis.xlsx Metrics','request instructions']},{'criterion_id':'scope','score':3,'rationale':'Memo retains bounded sample disclaimer','evidence':['memo.docx body']}]})
 queues={'author':[author],'solver':solve('r0',False)+solve('r1',True),'judge':[grade(False),grade(True)],'attributor':[json_reply({'attribution':[{'criterion_id':'threshold','labels':['Task'],'explanation':'Request lacks threshold required by rubric','evidence':['request.instructions lacks threshold; rubric threshold row']} ]})],'repairer':[repair],'repair_reviewer':[json_reply({'approved':True,'reasons':['Added missing source-supported threshold; all constraints preserved'],'weakened_criteria':[]})]}
 return manifest_path,manifest,expected,Transport(queues)
class FullProtocolAcceptance(unittest.TestCase):
 @classmethod
 def setUpClass(cls):
  (QA/'protocol_runs').mkdir(exist_ok=True);cls.root=Path(tempfile.mkdtemp(prefix='mock_http_',dir=QA/'protocol_runs'));mp,m,cls.expected,cls.http=setup_protocol();cls.config=cfg()
  with patch.dict(os.environ,{'WG_QA_ONLY_KEY':'QA_NOT_A_REAL_KEY_DO_NOT_USE'}),patch.object(rt.urllib.request,'build_opener',return_value=cls.http):
   cls.initial=wf.author_scenario({'profession':'Geoscientist','activity':'Analyze catalog facts'},m,mp.parent,cls.root/'authoring',cls.config,scenario_id='qa-protocol')
   cls.result=wf.evolve_scenario(cls.initial,cls.root/'workflow',cls.config)
  write_json(cls.root/'qa_transport_requests.json',cls.http.requests);write_json(cls.root/'QA_SIMULATION_ONLY.json',{'mode':'fully_mocked_http_protocol','genuine_model_calls':0,'training_eligible':False,'expected':cls.expected});write_json(QA/'latest_protocol.json',{'path':str(cls.root),'mode':'fully_mocked_http_protocol'})
 def test_complete_repair_sequence_uses_real_runtime_and_artifacts(self):
  self.assertEqual(self.result['status'],'accepted');self.assertEqual(len(self.result['rounds']),2)
  self.assertEqual([p['model'] for p in self.http.requests],['QA_MOCK_ONLY_author']+['QA_MOCK_ONLY_solver']*4+['QA_MOCK_ONLY_judge','QA_MOCK_ONLY_attributor','QA_MOCK_ONLY_repairer','QA_MOCK_ONLY_repair_reviewer']+['QA_MOCK_ONLY_solver']*4+['QA_MOCK_ONLY_judge'])
  self.assertTrue(all(not q for q in self.http.queues.values()))
 def test_solver_payload_never_has_private_rubric_or_repair_history(self):
  solver=[p for p in self.http.requests if p['model']=='QA_MOCK_ONLY_solver'];text=json.dumps(solver);self.assertNotIn(PRIVATE,text);self.assertNotIn(REPAIR,text);self.assertNotIn('private_reference',text);self.assertNotIn('"criterion_id"',text)
  for p in solver:
   for msg in p['messages']:
    if msg['role']=='user':self.assertNotIn('rubric',json.loads(msg['content']))
 def test_private_payloads_fresh_and_role_bounded(self):
  roles=[p for p in self.http.requests if p['model']!='QA_MOCK_ONLY_solver'];self.assertTrue(all(len(p['messages'])==2 for p in roles));self.assertTrue(any(PRIVATE in json.dumps(p) for p in roles if p['model']=='QA_MOCK_ONLY_judge'))
 def test_final_native_numeric_outputs_independently_recompute(self):
  from openpyxl import load_workbook
  from docx import Document
  scenario=self.root/'workflow'/self.result['accepted_scenario'];run=self.root/'workflow/round_001/executor';out=run/'workspace/output';book=load_workbook(out/'analysis.xlsx',data_only=True);values=dict(book['Metrics'].values);self.assertEqual(values['event_count'],self.expected['count']);self.assertEqual(values['threshold_count'],self.expected['threshold']);self.assertIn('not a complete census',' '.join(p.text for p in Document(out/'memo.docx').paragraphs));self.assertEqual(verify_bundle(scenario)['revision'],1);self.assertEqual(inspect_outputs(scenario,run)['missing'],[])
 def test_current_bundle_private_canary_excluded_from_public_files(self):
  current=self.root/'workflow'/self.result['accepted_scenario'];public=''.join(p.read_text(errors='ignore') for p in (current/'public').rglob('*') if p.is_file());self.assertNotIn(PRIVATE,public);self.assertNotIn(REPAIR,public)
 def test_malformed_audit_is_rejected_never_autoaccepted(self):
  scenario=self.root/'workflow'/self.result['accepted_scenario'];run=self.root/'workflow/round_001/executor'
  with patch.object(wf,'json_role',return_value={'result':{'scores':[]},'receipt':{}}),self.assertRaises(ForgeError):wf.audit_run(scenario,run,self.root/'malformed_audit',self.config)
  self.assertFalse((self.root/'malformed_audit/audit.json').exists())
 def test_fixture_never_automatically_becomes_training_data(self):
  with self.assertRaises(ForgeError):export_sft([],self.root/'no_qualified_export')
 def test_unreviewed_solver_trace_cannot_qualify(self):
  scenario=self.root/'workflow'/self.result['accepted_scenario'];run=self.root/'workflow/round_001/executor'
  with self.assertRaises(ForgeError):qualify(scenario,run,self.root/'workflow/round_001/audit/audit.json',self.root/'no_human_review.json',self.root/'must_not_qualify.json')
  self.assertFalse((self.root/'must_not_qualify.json').exists())
 def test_offline_fixture_mode_cannot_be_signed_as_live(self):
  scenario=self.root/'workflow'/self.result['accepted_scenario'];run=self.root/'signing_fixture';shutil.copytree(self.root/'workflow/round_001/executor',run);report=read_json(run/'run.json');report['mode']='offline_programmatic_control';write_json(run/'run.json',report)
  attestations={k:True for k in ['source_rights','semantic_correctness','privacy_safety','no_benchmark_overlap']}
  with self.assertRaises(ForgeError):sign_review(scenario,run,self.root/'workflow/round_001/audit/audit.json',self.root/'must_not_sign.json','QA TEST ONLY','QA_NONEXISTENT_REVIEW_KEY',attestations)
  self.assertFalse((self.root/'must_not_sign.json').exists())
 def test_cached_malformed_audit_is_rejected(self):
  current=self.root/'workflow'/self.result['accepted_scenario'];run=self.root/'workflow/round_001/executor';out=self.root/'malformed_cached_audit';out.mkdir();cached=read_json(self.root/'workflow/round_001/audit/audit.json');cached['scores']=[];write_json(out/'audit.json',cached)
  with patch.object(rt.urllib.request,'build_opener',side_effect=AssertionError('unexpected provider access')),self.assertRaises(ForgeError):wf.audit_run(current,run,out,self.config)
 def test_terminal_resume_revalidates_output_integrity(self):
  target=self.root/'tampered_terminal';shutil.copytree(self.root/'workflow',target)
  with (target/'round_001/executor/workspace/output/memo.docx').open('ab') as f:f.write(b'QA tamper after audit')
  with patch.object(rt.urllib.request,'build_opener',side_effect=AssertionError('unexpected provider access')),self.assertRaises(ForgeError):wf.evolve_scenario(self.initial,target,self.config)
 def test_resume_terminal_workflow_makes_zero_calls(self):
  with patch.object(rt.urllib.request,'build_opener',side_effect=AssertionError('unexpected provider access')):again=wf.evolve_scenario(self.initial,self.root/'workflow',self.config)
  self.assertEqual(again['status'],'accepted')
if __name__=='__main__':unittest.main(verbosity=2)
