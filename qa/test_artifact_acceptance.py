"""Independent artifact QA. No network, provider requests, or model code execution."""
from pathlib import Path
import tempfile,unittest,zipfile,hashlib,sys
from _qa_paths import REPO, QA_OUTPUT, TOKENIZER
sys.path.insert(0,str(REPO))
from wgforge.artifacts import write_artifact,extract_artifact
from wgforge.safeformula import evaluate_formula,FormulaError
class ArtifactAcceptance(unittest.TestCase):
 def setUp(self): self.t=tempfile.TemporaryDirectory();self.root=Path(self.t.name)
 def tearDown(self): self.t.cleanup()
 def spec(self): return {'title':'Public source analytical decision','text':'Evidence: 12 units at 2.5 per unit. Total 30.\nDecision: verify the source before approval.','tables':[{'name':'Metrics','headers':['item','units','rate','amount'],'rows':[['A',12,2.5,30],['B',8,3,24]]}]}
 def test_office_pdf_native_roundtrip(self):
  for ext in ['xlsx','docx','pptx','pdf','csv']:
   with self.subTest(ext=ext):
    p=self.root/f'out.{ext}';write_artifact(p,self.spec());r=extract_artifact(p)
    self.assertTrue(r['text'].strip());self.assertEqual(r['format'],ext)
    if ext in ['xlsx','docx','pptx']: self.assertTrue(zipfile.is_zipfile(p))
    elif ext=='pdf': self.assertTrue(p.read_bytes().startswith(b'%PDF-'))
 def test_generation_bytes_repeat(self):
  for ext in ['xlsx','docx','pptx','pdf','csv']:
   with self.subTest(ext=ext):
    p=self.root/f'a.{ext}';q=self.root/f'b.{ext}';write_artifact(p,self.spec());write_artifact(q,self.spec());self.assertEqual(p.read_bytes(),q.read_bytes())
 def test_invalid_table_dimensions_rejected(self):
  spec=self.spec();spec['tables'][0]['rows'][0].append('extra')
  with self.assertRaises(ValueError):write_artifact(self.root/'bad.xlsx',spec)
 def test_csv_cardinality_rejected(self):
  spec=self.spec();spec['tables'].append(spec['tables'][0])
  with self.assertRaises(ValueError):write_artifact(self.root/'bad.csv',spec)
 def test_fake_office_files_rejected(self):
  for ext in ['xlsx','docx','pptx','pdf']:
   with self.subTest(ext=ext):
    p=self.root/f'fake.{ext}';p.write_text('The requested artifact is complete.')
    with self.assertRaises(Exception):extract_artifact(p)
 def test_symlink_rejected(self):
  target=self.root/'target.txt';target.write_text('secret');link=self.root/'link.txt';link.symlink_to(target)
  with self.assertRaises(ValueError):extract_artifact(link)
 def test_macros_rejected(self):
  p=self.root/'macro.docx';write_artifact(p,self.spec())
  with zipfile.ZipFile(p,'a') as z:z.writestr('word/vbaProject.bin',b'fake')
  with self.assertRaises(ValueError):extract_artifact(p)
 def test_docx_header_and_footer_extracted(self):
  from docx import Document
  p=self.root/'document.docx';write_artifact(p,self.spec());d=Document(p);d.sections[0].header.paragraphs[0].text='HEADER_CANARY_2789';d.sections[0].footer.paragraphs[0].text='FOOTER_CANARY_7832';d.save(p)
  text=extract_artifact(p)['text'];self.assertIn('HEADER_CANARY_2789',text);self.assertIn('FOOTER_CANARY_7832',text)
 def test_pptx_notes_extracted(self):
  from pptx import Presentation
  p=self.root/'slides.pptx';write_artifact(p,self.spec());prs=Presentation(p);prs.slides[0].notes_slide.notes_text_frame.text='NOTE_CANARY_4289';prs.save(p);self.assertIn('NOTE_CANARY_4289',extract_artifact(p)['text'])
 def test_formula_independent_numeric_result(self):
  from openpyxl import Workbook
  w=Workbook();s=w.active;s.title='Data';s.append(['units','rate','value']);s.append([12,2.5,'=A2*B2']);s.append([8,3,'=A3*B3']);s.append(['total',None,'=SUM(C2:C3)'])
  p=self.root/'calc.xlsx';w.save(p);r=extract_artifact(p);self.assertEqual(r['formulas'][-1]['value'],12*2.5+8*3)
 def test_formula_cyclic_and_external_rejected(self):
  from openpyxl import Workbook
  w=Workbook();s=w.active;s['A1']='=A2';s['A2']='=A1'
  with self.assertRaises(FormulaError):evaluate_formula('=A1',w,s.title)
  for expr in ['=WEBSERVICE("https://example.org")','=__import__("os").system("echo bad")','=2^9999999','=1/0']:
   with self.subTest(expr=expr),self.assertRaises(FormulaError):evaluate_formula(expr)
 def test_image_only_docx_marked_unverified(self):
  from PIL import Image
  from docx import Document
  im=self.root/'diagram.png';Image.new('RGB',(20,20),'red').save(im);p=self.root/'picture.docx';d=Document();d.add_picture(str(im));d.save(p);self.assertTrue(extract_artifact(p)['unverified_visual_content'])
if __name__=='__main__':unittest.main(verbosity=2)
